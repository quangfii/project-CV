from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUTPUT = "Bao_cao_tong_quan_paper_SoilClass.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)
        borders.append(element)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def add_reference(doc, index, citation, label, url):
    paragraph = doc.add_paragraph(style="Body Text")
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.65)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.add_run(f"[{index}] {citation} ")
    add_hyperlink(paragraph, label, url)
    for run in paragraph.runs:
        run.font.size = Pt(9.2)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(1.9)
section.left_margin = Cm(2.35)
section.right_margin = Cm(2.15)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

body = styles["Body Text"]
body.font.name = "Arial"
body.font.size = Pt(10.5)
body._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
body.paragraph_format.space_after = Pt(6)
body.paragraph_format.line_spacing = 1.15
body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for style_name, size, before, after in (
    ("Title", 21, 0, 14),
    ("Subtitle", 11, 0, 16),
    ("Heading 1", 15, 14, 7),
    ("Heading 2", 12, 10, 5),
    ("Heading 3", 10.5, 8, 4),
):
    style = styles[style_name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

styles["Title"].font.bold = True
styles["Heading 1"].font.bold = True
styles["Heading 2"].font.bold = True
styles["Heading 3"].font.bold = True

title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for list_style in ("List Bullet", "List Number"):
    styles[list_style].font.name = "Arial"
    styles[list_style].font.size = Pt(10.5)
    styles[list_style]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles[list_style].paragraph_format.space_after = Pt(3)

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.add_run("SoilClass   |   ")
add_page_number(footer_paragraph)
for run in footer_paragraph.runs:
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Tổng quan nghiên cứu và lựa chọn phương pháp cho dự án SoilClass")

subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Phân loại ảnh đất bằng MobileNetV3 Small và transfer learning")

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run("Tài liệu phục vụ báo cáo và thuyết trình dự án")
run.font.name = "Arial"
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(90, 90, 90)

doc.add_paragraph()
add_body(
    doc,
    "Tài liệu này xác định vấn đề nghiên cứu, tổng hợp các nhóm giải pháp đã được công bố và giải thích vì sao dự án lựa chọn MobileNetV3-Small kết hợp transfer learning, data augmentation và class-weighted loss. Kết luận chính là giải pháp hiện tại phù hợp với mục tiêu xây dựng một hệ thống nhẹ, dễ triển khai và hoạt động tốt trên bộ SoilClass, nhưng kết quả kiểm thử ngoài miền cho thấy mô hình chưa đủ bằng chứng để sử dụng như một công cụ nhận dạng loại đất phổ quát.",
)

doc.add_paragraph(style="Heading 1").add_run("1 Topic")
add_body(
    doc,
    "Chủ đề của dự án là phân loại ảnh bề mặt đất bằng deep learning trong điều kiện dữ liệu hạn chế. Hệ thống nhận một ảnh RGB, chuẩn hóa ảnh về kích thước 224 x 224 pixel và dự đoán một trong năm lớp type_1 đến type_5 của bộ SoilClass. Mục tiêu kỹ thuật là đạt độ chính xác tốt với mô hình gọn nhẹ, có thể chạy trên CPU và tích hợp vào giao diện web để người dùng tải ảnh lên dự đoán.",
)
add_body(
    doc,
    "Tên đề tài đề xuất: Phân loại ảnh đất bằng mô hình MobileNetV3-Small sử dụng transfer learning và đánh giá khả năng tổng quát hóa trên dữ liệu ngoài miền.",
    bold_lead="Tên đề tài đề xuất:",
)

doc.add_paragraph(style="Heading 1").add_run("2 Problem")
add_body(
    doc,
    "Phân loại đất truyền thống thường cần kiến thức chuyên môn, phép đo vật lý hoặc phân tích trong phòng thí nghiệm. Phân loại từ ảnh có thể giảm thời gian và chi phí, nhưng đây là một bài toán khó vì màu sắc và kết cấu đất thay đổi theo độ ẩm, ánh sáng, thiết bị chụp, vị trí địa lý và cách chuẩn bị mẫu.",
)
add_body(doc, "Các vấn đề cụ thể của dự án gồm:")
for text in (
    "Bộ SoilClass chỉ có 552 ảnh và năm lớp, tương đối nhỏ đối với deep learning.",
    "Các lớp được đặt tên type_1 đến type_5, không có bảng ánh xạ sang tên đất vật lý như sand, clay hoặc loam.",
    "Số ảnh giữa các lớp không hoàn toàn cân bằng, với tổng số lần lượt là 109, 96, 137, 107 và 103 ảnh.",
    "Ảnh train, validation và test đến từ cùng một nguồn nên độ chính xác nội miền có thể cao hơn hiệu năng ngoài thực tế.",
    "Mô hình có nguy cơ học đặc điểm của camera, ánh sáng hoặc bối cảnh thay vì đặc tính đất có khả năng tổng quát hóa.",
    "Giao diện hiện luôn ép ảnh đầu vào vào một trong năm lớp và chưa có cơ chế từ chối ảnh ngoài phân phối.",
):
    add_bullet(doc, text)

doc.add_paragraph(style="Heading 1").add_run("3 Các solution đã có")

doc.add_paragraph(style="Heading 2").add_run("3.1 Đặc trưng thủ công và machine learning")
add_body(
    doc,
    "Các phương pháp truyền thống trích xuất màu sắc, histogram, texture, HOG, LBP hoặc GLCM rồi dùng SVM, k-NN, Random Forest hay Decision Tree. Ưu điểm là dễ giải thích và có thể hoạt động với ít dữ liệu. Nhược điểm là phụ thuộc mạnh vào cách thiết kế đặc trưng và khó thích nghi với sự thay đổi ánh sáng, độ ẩm và nền ảnh.",
)

doc.add_paragraph(style="Heading 2").add_run("3.2 CNN tự thiết kế")
add_body(
    doc,
    "Light-SoilNet là ví dụ tiêu biểu cho CNN nhẹ được thiết kế riêng cho ảnh đất. Nghiên cứu thu thập 392 mẫu đất, xác định nhãn bằng sieve và hydrometer, tiền xử lý RGB, kênh V của HSV và adaptive histogram, sau đó phân loại năm loại đất. Mô hình đạt accuracy 97,2%. Cách tiếp cận này gần với dự án về bài toán, quy mô dữ liệu và yêu cầu mô hình nhẹ, nhưng cần thiết kế và kiểm chứng kiến trúc riêng.",
)

doc.add_paragraph(style="Heading 2").add_run("3.3 Transfer learning với CNN pretrained")
add_body(
    doc,
    "Transfer learning sử dụng backbone đã học trên ImageNet, sau đó thay và huấn luyện classifier cho bài toán mới. SoilNet-XAI so sánh VGG16, VGG19, ResNet50, InceptionV3, DenseNet121 và EfficientNet trên 368 ảnh đất. DenseNet121 đạt accuracy 98,08%, precision 98,16%, recall 98,08% và F1-score 98,06%. Đây là nhóm giải pháp gần nhất với mã nguồn hiện tại của dự án.",
)

doc.add_paragraph(style="Heading 2").add_run("3.4 Tiền xử lý và tách nền")
add_body(
    doc,
    "Một nghiên cứu trên Computers and Electronics in Agriculture sử dụng YOLOv5 để loại nền trước khi phân loại 720 ảnh thuộc sáu loại đất. DenseNet121 đạt 97,22%, Xception đạt 93,52%, trong khi MobileNetV2 đạt 72,22% nhưng có thời gian huấn luyện nhanh nhất và mức dùng bộ nhớ thấp nhất. Giải pháp tách nền có thể giảm việc mô hình học các chi tiết không thuộc đất, đổi lại pipeline phức tạp hơn và cần thêm mô hình detection.",
)

doc.add_paragraph(style="Heading 2").add_run("3.5 Mô hình hybrid và ensemble")
add_body(
    doc,
    "Các mô hình hybrid kết hợp nhiều backbone để thu được đặc trưng bổ sung. Paper Scientific Reports năm 2026 ghép ResNet50 và MobileNetV2 bằng projection và feature concatenation, đạt accuracy 73,65% trên Soil_Data_V3 gồm 5.254 ảnh và tám lớp. Một nghiên cứu khác trên Engineering Applications of Artificial Intelligence kết hợp CNN, Vision Transformer, CycleGAN và fuzzy logic. Các giải pháp này có tiềm năng cao nhưng nặng hơn, khó huấn luyện và khó triển khai hơn MobileNetV3-Small.",
)

doc.add_paragraph(style="Heading 2").add_run("3.6 Explainable AI và phát hiện dữ liệu ngoài miền")
add_body(
    doc,
    "Grad-CAM và LIME được SoilNet-XAI sử dụng để kiểm tra vùng ảnh ảnh hưởng đến dự đoán. Đây là bước quan trọng để xác định mô hình tập trung vào kết cấu đất hay vào phông nền. Với hệ thống triển khai thực tế, explainability nên đi cùng confidence calibration và cơ chế từ chối ảnh ngoài phân phối.",
)

heading_four = doc.add_paragraph(style="Heading 1")
heading_four.paragraph_format.page_break_before = True
heading_four.add_run("4 Đánh giá vấn đề và giải pháp")

table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(3.05), Cm(3.1), Cm(3.4), Cm(3.2), Cm(3.45)]
headers = ["Nhóm giải pháp", "Ưu điểm", "Hạn chế", "Mức phù hợp", "Vai trò trong báo cáo"]
for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
    cell.width = width
    cell.text = header
    set_cell_shading(cell, "1F4E78")
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.5)

comparison_rows = [
    ("Đặc trưng thủ công + ML", "Nhẹ, dễ giải thích, cần ít dữ liệu", "Phụ thuộc feature engineering", "Trung bình", "Baseline truyền thống"),
    ("CNN tự thiết kế", "Tối ưu cho bài toán", "Cần thiết kế và ablation kỹ", "Cao", "So sánh với Light-SoilNet"),
    ("Transfer learning", "Hiệu quả với dataset nhỏ, triển khai nhanh", "Có thể lệch miền so với ImageNet", "Rất cao", "Nền tảng phương pháp đã chọn"),
    ("Tách nền + CNN", "Giảm nhiễu bối cảnh", "Tăng chi phí và độ phức tạp", "Cao", "Hướng cải tiến"),
    ("Hybrid hoặc ensemble", "Đặc trưng phong phú", "Nặng, khó triển khai trên CPU", "Trung bình", "Related work và hướng tương lai"),
    ("XAI và OOD", "Tăng độ tin cậy và minh bạch", "Không tự cải thiện accuracy", "Rất cao", "Bổ sung cần thiết trước triển khai"),
]
for row_index, values in enumerate(comparison_rows, start=1):
    cells = table.add_row().cells
    for col_index, (cell, value, width) in enumerate(zip(cells, values, widths)):
        cell.width = width
        cell.text = value
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_index % 2 == 0:
            set_cell_shading(cell, "EAF2F8")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index < 3 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.2)
set_table_borders(table)

add_body(
    doc,
    "Đánh giá tổng hợp cho thấy transfer learning với backbone nhẹ là lựa chọn cân bằng nhất cho phạm vi dự án. CNN tự thiết kế có thể đạt kết quả cao nhưng cần nhiều thí nghiệm kiến trúc. DenseNet hoặc hybrid có thể cải thiện khả năng biểu diễn nhưng tăng chi phí tính toán. Tách nền, Grad-CAM và phát hiện ảnh ngoài miền là các bổ sung hợp lý hơn cho giai đoạn tiếp theo.",
)

doc.add_paragraph(style="Heading 1").add_run("5 Bài survey trung tâm")
doc.add_paragraph(style="Heading 2").add_run("Soil image classification and segmentation A survey from deep learning multimodal data and hybrid models")
add_body(
    doc,
    "Bài survey được công bố trên GeoAI, Elsevier, năm 2026. Công trình tổng hợp tiến triển của phân tích ảnh đất từ đặc trưng thủ công và bộ phân loại thống kê đến CNN, transformer, self-supervised learning, transfer learning, mô hình hybrid và dữ liệu đa phương thức. Survey phân biệt ba nhóm nhiệm vụ chính là phân loại tổng quát, object detection và segmentation theo pixel; đồng thời liên hệ chúng với nhận dạng loại đất, thành phần khoáng, tính chất đất và phân tích hình thái hạt.",
)
add_body(
    doc,
    "Đóng góp quan trọng của survey là chỉ ra rằng transfer learning và self-supervised learning phù hợp với bối cảnh thiếu dữ liệu gán nhãn; segmentation hỗ trợ phân tích kích thước và phân bố hạt; còn hybrid và multimodal learning là hướng phát triển để tăng độ bền vững. Những thách thức xuyên suốt gồm dataset nhỏ, mất cân bằng lớp, khác biệt thiết bị và môi trường, thiếu chuẩn đánh giá chung, khả năng giải thích và tổng quát hóa ngoài miền.",
)
add_body(
    doc,
    "Liên hệ với dự án: MobileNetV3-Small thuộc nhóm transfer learning với CNN; augmentation và class-weighted loss giải quyết một phần thiếu dữ liệu và mất cân bằng; external test phản ánh trực tiếp thách thức domain shift mà survey nêu. Dự án chưa có segmentation, multimodal data hoặc XAI, nên đây là các giới hạn và hướng phát triển hợp lý.",
)
survey_link = doc.add_paragraph(style="Body Text")
survey_link.add_run("Liên kết survey: ").bold = True
add_hyperlink(survey_link, "ScienceDirect và DOI 10.1016/j.geoai.2026.100053", "https://doi.org/10.1016/j.geoai.2026.100053")

doc.add_paragraph(style="Heading 1").add_run("6 Các paper liên quan và so sánh với giải pháp đã chọn")

papers = [
    (
        "6.1 Searching for MobileNetV3",
        "Howard và cộng sự giới thiệu MobileNetV3, được tối ưu cho CPU di động bằng hardware-aware neural architecture search, NetAdapt và các cải tiến kiến trúc. Paper cung cấp MobileNetV3-Large và MobileNetV3-Small cho các mức tài nguyên khác nhau.",
        "Đây là nền tảng trực tiếp của backbone trong dự án. MobileNetV3-Small phù hợp vì file mô hình gọn, inference nhanh và có thể triển khai trên CPU. Dự án sử dụng bản pretrained ImageNet của TorchVision và thay classifier theo năm lớp SoilClass.",
        "https://openaccess.thecvf.com/content_ICCV_2019/html/Howard_Searching_for_MobileNetV3_ICCV_2019_paper.html",
    ),
    (
        "6.2 Smart soil image classification system using lightweight convolutional neural network",
        "Pandiri, Murugan và Goel đề xuất Light-SoilNet cho 392 mẫu thuộc năm loại đất, kết hợp thông tin RGB, HSV và adaptive histogram. Mô hình đạt accuracy 97,2% và được thiết kế để xử lý dữ liệu mất cân bằng.",
        "Light-SoilNet gần nhất với mục tiêu của dự án, nhưng yêu cầu kiến trúc và preprocessing riêng. MobileNetV3-Small được chọn vì có backbone đã được kiểm chứng, trọng số pretrained và hệ sinh thái triển khai tốt; cách này giảm rủi ro khi chỉ có 552 ảnh.",
        "https://www.sciencedirect.com/science/article/pii/S0957417423026878",
    ),
    (
        "6.3 SoilNet XAI",
        "Paper IEEE đánh giá nhiều mạng pretrained trên 368 ảnh thuộc bốn lớp. DenseNet121 đạt accuracy 98,08% và F1-score 98,06%. Grad-CAM và LIME được dùng để kiểm tra vùng ảnh quyết định dự đoán.",
        "Giải pháp của dự án nhẹ hơn DenseNet121 và có accuracy nội miền tương đương, nhưng chưa có explainability. SoilNet-XAI hỗ trợ đề xuất bổ sung Grad-CAM để xác minh mô hình nhìn vào texture của đất thay vì nền ảnh.",
        "https://ieeexplore.ieee.org/document/11491058/",
    ),
    (
        "6.4 Performance comparison of deep and transfer learning models for smart soil texture classification",
        "Nghiên cứu dùng YOLOv5 để loại nền rồi đánh giá chín mô hình transfer learning trên 720 ảnh và sáu loại đất. DenseNet121 đạt 97,22%, Xception đạt 93,52%, MobileNetV2 đạt 72,22%; MobileNetV2 nhanh nhất và dùng ít bộ nhớ nhất.",
        "Paper cho thấy đánh đổi giữa accuracy và tài nguyên. Dự án ưu tiên MobileNetV3-Small vì cần triển khai nhẹ, nhưng tách nền bằng detector là hướng cải tiến nếu kiểm thử cho thấy mô hình đang dựa vào bối cảnh.",
        "https://www.sciencedirect.com/science/article/pii/S0168169925008282",
    ),
    (
        "6.5 Hybrid deep learning model for automated soil classification and crop recommendation",
        "Paper Scientific Reports ghép ResNet50 và MobileNetV2. Hai vector đặc trưng được chiếu về 512 chiều, nối thành vector 1.024 chiều và phân loại bằng các dense layer. Hybrid đạt accuracy 73,65%, macro F1 70,97% trên 526 ảnh test của Soil_Data_V3.",
        "Hybrid tận dụng đặc trưng bổ sung nhưng tốn tài nguyên hơn. Với dataset chỉ 552 ảnh và môi trường CPU, dự án chọn một backbone nhẹ để giảm độ phức tạp và nguy cơ overfitting. Hybrid phù hợp làm hướng nghiên cứu sau khi có dữ liệu lớn hơn.",
        "https://www.nature.com/articles/s41598-026-66201-0",
    ),
    (
        "6.6 Advanced AI framework integrating CNNs and Vision Transformers",
        "Nghiên cứu EAAI kết hợp CNN, Vision Transformer, CycleGAN và fuzzy logic cho phân loại đất và gợi ý cây trồng. Dataset gốc gồm 1.189 ảnh và bảy loại đất; accuracy tối đa trên dữ liệu gốc là 89,32%, tăng lên 91,01% sau augmentation theo báo cáo của tác giả.",
        "Giải pháp đầy đủ hơn nhưng phức tạp và phụ thuộc dữ liệu sinh. Dự án đã dùng chính 1.189 ảnh gốc của nghiên cứu này để kiểm thử zero-shot mà không huấn luyện lại, qua đó đánh giá domain shift thay vì chỉ trình bày accuracy nội miền.",
        "https://doi.org/10.1016/j.engappai.2025.111425",
    ),
]

for heading, description, comparison, url in papers:
    doc.add_paragraph(style="Heading 2").add_run(heading)
    add_body(doc, description)
    add_body(doc, "So sánh với dự án: " + comparison, bold_lead="So sánh với dự án:")
    link_paragraph = doc.add_paragraph(style="Body Text")
    link_paragraph.add_run("Paper: ").bold = True
    add_hyperlink(link_paragraph, url, url)

doc.add_paragraph(style="Heading 1").add_run("7 Phương pháp dự án đã chọn")
add_body(
    doc,
    "Phương pháp hiện tại sử dụng MobileNetV3-Small pretrained trên ImageNet. Ảnh train được RandomResizedCrop về 224 x 224, lật ngang, xoay nhẹ và điều chỉnh brightness, contrast, saturation. Ảnh validation và test chỉ resize, center crop và normalize. Feature extractor được đóng băng, classifier được thay bằng đầu ra năm lớp và huấn luyện bằng AdamW.",
)
add_body(doc, "Các thành phần chính:")
for text in (
    "Transfer learning để tận dụng đặc trưng thị giác đã học từ ImageNet.",
    "MobileNetV3-Small để giảm kích thước mô hình và chi phí inference.",
    "Data augmentation để tăng biến thiên của dataset nhỏ.",
    "Class-weighted cross-entropy để giảm ảnh hưởng của chênh lệch số lượng lớp.",
    "Learning-rate scheduler và early stopping dựa trên validation accuracy.",
    "Đánh giá bằng accuracy, precision, recall, F1-score và confusion matrix.",
    "Giao diện Gradio cho phép tải ảnh và hiển thị xác suất năm lớp.",
    "External test chạy inference trên dữ liệu paper khác mà không cập nhật trọng số.",
):
    add_bullet(doc, text)

doc.add_paragraph(style="Heading 2").add_run("7.1 Kết quả nội miền")
add_body(
    doc,
    "Mô hình đạt validation accuracy 98,2% và test accuracy 98,15%, tương ứng 106 dự đoán đúng trên 108 ảnh test. Type_1, type_2 và type_3 đạt precision và recall 100%; type_4 có recall 90,48%; type_5 có precision 90,91%.",
)

doc.add_paragraph(style="Heading 2").add_run("7.2 Kết quả ngoài miền")
add_body(
    doc,
    "Mô hình được kiểm thử không huấn luyện lại trên 1.189 ảnh thuộc bảy loại đất từ dataset của paper EAAI. Hệ thống xử lý thành công toàn bộ ảnh, nhưng confidence trung bình chỉ đạt 63,75% và 47,1% ảnh có confidence dưới 60%. Các lớp vật lý bên ngoài bị phân tán qua nhiều type, cho thấy khả năng tổng quát hóa còn hạn chế.",
)

doc.add_paragraph(style="Heading 1").add_run("8 Vì sao chọn giải pháp này")
reasons = [
    ("Phù hợp với dữ liệu nhỏ", "Transfer learning giảm nhu cầu học đặc trưng từ đầu khi chỉ có 389 ảnh train."),
    ("Phù hợp với phần cứng", "MobileNetV3-Small có kích thước nhỏ và inference được trên CPU, thuận lợi cho demo và triển khai web."),
    ("Ít rủi ro hơn CNN tự thiết kế", "Backbone đã được kiểm chứng và có trọng số ImageNet, trong khi CNN mới cần nhiều ablation và dữ liệu hơn."),
    ("Đơn giản hơn hybrid", "Một backbone giúp quy trình huấn luyện, lưu mô hình và triển khai dễ tái lập hơn ResNet-MobileNet fusion hoặc CNN-ViT ensemble."),
    ("Có biện pháp cho mất cân bằng", "Class weights bổ sung trực tiếp vào loss thay vì phải nhân bản ảnh thủ công."),
    ("Có thể mở rộng", "Pipeline cho phép thay số lớp, bổ sung Grad-CAM, background removal, confidence threshold hoặc fine-tuning khi có thêm dữ liệu."),
]
for lead, explanation in reasons:
    add_body(doc, f"{lead}: {explanation}", bold_lead=f"{lead}:")

add_body(
    doc,
    "Do đó, lựa chọn MobileNetV3-Small không nhằm tuyên bố đây là kiến trúc có accuracy cao nhất trong mọi tình huống. Đây là lựa chọn cân bằng giữa độ chính xác nội miền, kích thước mô hình, thời gian huấn luyện, khả năng chạy CPU và mức độ phức tạp phù hợp với phạm vi đồ án.",
)

doc.add_paragraph(style="Heading 1").add_run("9 Kết luận và hướng phát triển")
add_body(
    doc,
    "Dự án đã hoàn thành pipeline deep learning từ chuẩn bị dữ liệu, huấn luyện, đánh giá, lưu mô hình, dự đoán ảnh mới đến giao diện web. Kết quả nội miền cho thấy MobileNetV3-Small phù hợp với SoilClass. Tuy nhiên, external test chứng minh rằng mô hình chưa tổng quát hóa tốt sang nguồn ảnh khác và chưa thể quy đổi type_1 đến type_5 thành tên đất vật lý.",
)
add_body(doc, "Các bước ưu tiên cho nghiên cứu tiếp theo:")
for text in (
    "Xác minh ý nghĩa vật lý của type_1 đến type_5 hoặc xây dựng lại dataset với nhãn chuyên gia.",
    "Bổ sung ảnh từ nhiều thiết bị, địa điểm, điều kiện sáng và độ ẩm.",
    "Sử dụng group split theo phiên chụp hoặc mẫu đất để tránh ảnh gần giống xuất hiện ở cả train và test.",
    "Thêm Grad-CAM để kiểm tra vùng mô hình sử dụng khi dự đoán.",
    "Thêm confidence calibration và cơ chế từ chối ảnh ngoài phân phối.",
    "So sánh MobileNetV3 với DenseNet121, EfficientNet-B0 và một baseline đặc trưng thủ công trên cùng split.",
):
    add_bullet(doc, text)

doc.add_paragraph(style="Heading 1").add_run("10 Tài liệu tham khảo")
references = [
    (
        "SoilClass: A Novel Image Dataset for Soil Classification. Mendeley Data, Version 1, 2026. DOI: 10.17632/6mx34nshm4.1.",
        "Mendeley Data",
        "https://data.mendeley.com/datasets/6mx34nshm4/1",
    ),
    (
        "Howard, A. et al. Searching for MobileNetV3. Proceedings of ICCV, 2019, pp. 1314-1324.",
        "CVF Open Access",
        "https://openaccess.thecvf.com/content_ICCV_2019/html/Howard_Searching_for_MobileNetV3_ICCV_2019_paper.html",
    ),
    (
        "Pandiri, D. N. K., Murugan, R., and Goel, T. Smart soil image classification system using lightweight convolutional neural network. Expert Systems with Applications, 238, 122185, 2024.",
        "DOI",
        "https://doi.org/10.1016/j.eswa.2023.122185",
    ),
    (
        "SoilNet-XAI: A Deep Learning Approach with Explainability for Soil Image Classification. 28th ICCIT, IEEE, 2025. DOI: 10.1109/ICCIT68739.2025.11491058.",
        "IEEE Xplore",
        "https://ieeexplore.ieee.org/document/11491058/",
    ),
    (
        "Cetin, N. et al. Performance comparison of deep and transfer learning models for smart soil texture classification. Computers and Electronics in Agriculture, 237, 110722, 2025.",
        "ScienceDirect",
        "https://www.sciencedirect.com/science/article/pii/S0168169925008282",
    ),
    (
        "Sreelatha, A. et al. Hybrid deep learning model for automated soil classification and crop recommendation to enhance agricultural productivity. Scientific Reports, 2026.",
        "Nature",
        "https://www.nature.com/articles/s41598-026-66201-0",
    ),
    (
        "Sheth, F. et al. An advanced artificial intelligence framework integrating ensembled convolutional neural networks and Vision Transformers for precise soil classification with adaptive fuzzy logic-based crop recommendations. Engineering Applications of Artificial Intelligence, 158, 111425, 2025.",
        "DOI",
        "https://doi.org/10.1016/j.engappai.2025.111425",
    ),
    (
        "Soil image classification and segmentation: A survey from deep learning, multimodal data and hybrid models. GeoAI, 8, 100053, 2026.",
        "DOI",
        "https://doi.org/10.1016/j.geoai.2026.100053",
    ),
]
for index, (citation, label, url) in enumerate(references, start=1):
    add_reference(doc, index, citation, label, url)

doc.core_properties.title = "Tổng quan nghiên cứu và lựa chọn phương pháp cho dự án SoilClass"
doc.core_properties.subject = "Phân loại ảnh đất bằng MobileNetV3-Small và transfer learning"
doc.core_properties.keywords = "Soil classification, MobileNetV3, transfer learning, deep learning"
doc.core_properties.author = "Nhóm dự án SoilClass"
doc.save(OUTPUT)
print(OUTPUT)
