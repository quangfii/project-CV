from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = "Bao_cao_tong_quan_paper_SoilClass.docx"


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    pr.append(node)


def margins(cell, value=100):
    pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for name in ("top", "start", "bottom", "end"):
        item = OxmlElement(f"w:{name}")
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
        tc_mar.append(item)
    pr.append(tc_mar)


def borders(table):
    pr = table._tbl.tblPr
    node = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), "D9D9D9")
        node.append(edge)
    pr.append(node)


def hyperlink(paragraph, label, url):
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def body(doc, text, lead=None):
    p = doc.add_paragraph(style="Body Text")
    if lead and text.startswith(lead):
        p.add_run(lead).bold = True
        p.add_run(text[len(lead):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def paper_link(doc, label, url):
    p = doc.add_paragraph(style="Body Text")
    p.add_run("Liên kết: ").bold = True
    hyperlink(p, label, url)


def reference(doc, number, text, label, url):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.65)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    p.add_run(f"[{number}] {text} ")
    hyperlink(p, label, url)
    for run in p.runs:
        run.font.size = Pt(9.2)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.35)
section.right_margin = Cm(2.15)

for style_name in ("Normal", "Body Text", "List Bullet"):
    style = doc.styles[style_name]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

doc.styles["Body Text"].paragraph_format.space_after = Pt(6)
doc.styles["Body Text"].paragraph_format.line_spacing = 1.15
doc.styles["Body Text"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
doc.styles["List Bullet"].paragraph_format.space_after = Pt(3)

for name, size, before, after in (
    ("Title", 21, 0, 14),
    ("Subtitle", 11, 0, 16),
    ("Heading 1", 15, 14, 7),
    ("Heading 2", 12, 10, 5),
):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = name != "Subtitle"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title_pr = doc.styles["Title"]._element.get_or_add_pPr()
title_border = title_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_pr.remove(title_border)

footer = section.footer.paragraphs[0]
footer.add_run("Báo cáo tuần 1   |   ")
page_number(footer)
for run in footer.runs:
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Tổng quan bài toán phân loại ảnh đất")

subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Báo cáo đọc paper tuần đầu và định hướng lựa chọn phương pháp")

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = intro.add_run("Phạm vi tài liệu chỉ bao gồm kiến thức và nhận định rút ra từ các công trình đã công bố")
run.font.name = "Arial"
run.font.size = Pt(9.5)
run.font.color.rgb = RGBColor(90, 90, 90)

body(
    doc,
    "Báo cáo tuần đầu tập trung hiểu bài toán phân loại ảnh đất và khảo sát các phương pháp đã được công bố. Qua các paper, transfer learning với một CNN nhẹ được xem là hướng phù hợp cho bài toán có số lượng ảnh hạn chế và có nhu cầu triển khai thực tế. MobileNetV3-Small được chọn làm phương án dự kiến vì có cơ sở kiến trúc rõ ràng, chi phí tính toán thấp và phù hợp với xu hướng lightweight soil classification trong literature. Báo cáo này chưa trình bày triển khai, mã nguồn hoặc kết quả thực nghiệm của dự án.",
)

doc.add_paragraph(style="Heading 1").add_run("1 Topic")
body(
    doc,
    "Chủ đề nghiên cứu là phân loại loại đất từ ảnh RGB bằng deep learning. Hệ thống dự kiến nhận ảnh chụp bề mặt hoặc mẫu đất và trả về nhãn loại đất. Đây là bài toán computer vision thuộc nhóm supervised multiclass image classification.",
)
body(
    doc,
    "Ý nghĩa của bài toán nằm ở khả năng hỗ trợ nhận dạng đất nhanh và chi phí thấp hơn so với những quy trình chỉ dựa vào quan sát chuyên gia hoặc phân tích phòng thí nghiệm. Tuy nhiên, dự đoán từ ảnh chỉ khai thác đặc điểm thị giác; nó không thay thế trực tiếp các phép đo thành phần hóa học, độ pH hoặc kích thước hạt nếu dataset không cung cấp các nhãn đó.",
)

doc.add_paragraph(style="Heading 1").add_run("2 Problem formulation")

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(3.0), Cm(4.0), Cm(8.5)]
headers = ["Thành phần", "Định nghĩa", "Diễn giải"]
for cell, text, width in zip(table.rows[0].cells, headers, widths):
    cell.width = width
    cell.text = text
    shade(cell, "1F4E78")
    margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9)

problem_rows = [
    ("Input", "Ảnh RGB của đất", "Ảnh có thể khác nhau về độ phân giải, góc chụp, ánh sáng, độ ẩm, thiết bị và phông nền."),
    ("Output", "Một nhãn lớp", "Ví dụ sand, clay, loam hoặc các nhãn lớp do dataset định nghĩa."),
    ("Task", "Multiclass classification", "Mỗi ảnh được gán vào đúng một trong K lớp đất."),
    ("Learning setting", "Supervised learning", "Mô hình học từ các cặp ảnh và nhãn đã được cung cấp."),
    ("Mục tiêu", "Khả năng phân biệt và tổng quát hóa", "Không chỉ đúng trên ảnh cùng nguồn mà cần ổn định trước thay đổi điều kiện chụp."),
]
for row_index, values in enumerate(problem_rows, start=1):
    cells = table.add_row().cells
    for col_index, (cell, value, width) in enumerate(zip(cells, values, widths)):
        cell.width = width
        cell.text = value
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_index % 2 == 0:
            shade(cell, "EAF2F8")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index < 2 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(9)
borders(table)

doc.add_paragraph(style="Heading 2").add_run("2.1 Những khó khăn chính")
for text in (
    "Dataset ảnh đất thường nhỏ hoặc mất cân bằng giữa các lớp.",
    "Các lớp đất có thể gần nhau về màu sắc và kết cấu nên dễ bị nhầm lẫn.",
    "Ánh sáng, độ ẩm và camera làm thay đổi mạnh biểu hiện của cùng một loại đất.",
    "Phông nền có thể tạo shortcut khiến mô hình không học đúng đặc điểm bề mặt đất.",
    "Cách định nghĩa lớp khác nhau giữa các dataset, gây khó khăn cho đánh giá chéo.",
    "Accuracy nội bộ cao chưa chắc phản ánh khả năng tổng quát hóa sang địa điểm hoặc thiết bị khác.",
):
    p = bullet(doc, text)
    p.paragraph_format.space_after = Pt(0)

doc.add_paragraph(style="Heading 1").add_run("3 Evaluation metrics")
body(
    doc,
    "Các paper thường sử dụng nhiều metric vì accuracy đơn lẻ có thể che khuất hiệu năng kém ở lớp ít ảnh. Với bài toán multiclass, metric có thể được tính theo từng lớp rồi tổng hợp bằng macro average hoặc weighted average.",
)

metrics = doc.add_table(rows=1, cols=3)
metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
metrics.autofit = False
metric_widths = [Cm(3.1), Cm(5.2), Cm(7.2)]
for cell, text, width in zip(metrics.rows[0].cells, ["Metric", "Ý nghĩa", "Khi nào quan trọng"], metric_widths):
    cell.width = width
    cell.text = text
    shade(cell, "1F4E78")
    margins(cell)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9)

metric_rows = [
    ("Accuracy", "Tỷ lệ toàn bộ dự đoán đúng", "Phù hợp khi các lớp tương đối cân bằng."),
    ("Precision", "Trong các ảnh dự đoán là một lớp, bao nhiêu ảnh đúng", "Quan trọng khi false positive gây ảnh hưởng."),
    ("Recall", "Trong các ảnh thực sự thuộc một lớp, mô hình nhận ra bao nhiêu", "Quan trọng với lớp không được phép bỏ sót."),
    ("F1 score", "Trung bình điều hòa của precision và recall", "Hữu ích khi dữ liệu mất cân bằng."),
    ("Confusion matrix", "Số dự đoán đúng và nhầm giữa từng cặp lớp", "Giúp xác định các lớp có hình thái gần nhau."),
    ("ROC AUC", "Khả năng xếp hạng lớp theo nhiều threshold", "Dùng trong một số paper để phân tích khả năng phân biệt từng lớp."),
]
for row_index, values in enumerate(metric_rows, start=1):
    cells = metrics.add_row().cells
    for col_index, (cell, value, width) in enumerate(zip(cells, values, metric_widths)):
        cell.width = width
        cell.text = value
        margins(cell)
        if row_index % 2 == 0:
            shade(cell, "EAF2F8")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(9)
borders(metrics)

doc.add_paragraph(style="Heading 1").add_run("4 Bài survey trung tâm")
doc.add_paragraph(style="Heading 2").add_run("Soil image classification and segmentation A survey from deep learning multimodal data and hybrid models")
body(
    doc,
    "Survey trên GeoAI năm 2026 tổng hợp quá trình phát triển của phân tích ảnh đất từ đặc trưng thủ công và bộ phân loại thống kê đến CNN, transformer, transfer learning, self-supervised learning, mô hình hybrid và dữ liệu đa phương thức. Công trình phân biệt ba nhóm nhiệm vụ chính gồm classification, object detection và pixel-wise segmentation.",
)
body(
    doc,
    "Đối với classification, survey cho thấy CNN có khả năng học trực tiếp đặc trưng màu sắc, texture và hình thái. Transfer learning và self-supervised learning được xem là hướng quan trọng khi dữ liệu gán nhãn hạn chế. Với các bài toán cần thông tin chi tiết về hạt đất, segmentation được dùng để phân tích kích thước, hình dạng và phân bố không gian.",
)
body(
    doc,
    "Survey cũng nhấn mạnh các vấn đề còn mở: dataset nhỏ và thiếu chuẩn hóa, class imbalance, khác biệt môi trường thu thập, thiếu khả năng giải thích, khó tổng quát hóa và thiếu kiểm thử trên nguồn dữ liệu độc lập. Hướng phát triển được đề cập gồm multimodal learning, mô hình hybrid, self-supervised pretraining và đánh giá chặt chẽ hơn ngoài miền dữ liệu huấn luyện.",
)
paper_link(doc, "DOI 10.1016/j.geoai.2026.100053", "https://doi.org/10.1016/j.geoai.2026.100053")

doc.add_page_break()
doc.add_paragraph(style="Heading 1").add_run("5 Các phương pháp từ paper liên quan")

papers = [
    (
        "5.1 Light SoilNet",
        "Pandiri, Murugan và Goel xây dựng Light-SoilNet để phân loại sand, clay, loam, loamy sand và sandy loam. Trên 392 mẫu được xác định bằng sieve và hydrometer, nghiên cứu kết hợp RGB, kênh V của HSV và adaptive histogram để làm rõ texture. Mô hình đạt accuracy 97,2% và được thiết kế cho dữ liệu mất cân bằng.",
        "Điểm rút ra: một CNN nhẹ có thể đạt hiệu năng tốt trong bài toán ảnh đất quy mô nhỏ, nhưng chất lượng nhãn và quy trình thu thập mẫu đóng vai trò quan trọng.",
        "ScienceDirect",
        "https://www.sciencedirect.com/science/article/pii/S0957417423026878",
    ),
    (
        "5.2 SoilNet XAI",
        "SoilNet-XAI đánh giá các kiến trúc pretrained như VGG16, VGG19, ResNet50, InceptionV3, DenseNet121 và EfficientNet trên 368 ảnh thuộc bốn loại đất. DenseNet121 đạt accuracy 98,08%, precision 98,16%, recall 98,08% và F1-score 98,06%. Grad-CAM và LIME được dùng để giải thích vùng ảnh ảnh hưởng đến quyết định.",
        "Điểm rút ra: transfer learning phù hợp với dataset nhỏ; accuracy cần đi cùng F1, confusion matrix và phương pháp giải thích mô hình.",
        "IEEE Xplore",
        "https://ieeexplore.ieee.org/document/11491058/",
    ),
    (
        "5.3 Transfer learning kết hợp tách nền",
        "Cetin và cộng sự sử dụng YOLOv5 để loại nền trước khi so sánh chín mô hình transfer learning trên 720 ảnh thuộc sáu loại đất. DenseNet121 đạt 97,22%, Xception đạt 93,52% và MobileNetV2 đạt 72,22%. MobileNetV2 có thời gian huấn luyện nhanh và mức sử dụng bộ nhớ thấp nhất trong so sánh được báo cáo.",
        "Điểm rút ra: tách nền có thể giúp mô hình tập trung vào đất; đồng thời tồn tại đánh đổi giữa độ chính xác, tốc độ và tài nguyên tính toán.",
        "ScienceDirect",
        "https://www.sciencedirect.com/science/article/pii/S0168169925008282",
    ),
    (
        "5.4 Hybrid ResNet50 và MobileNetV2",
        "Sreelatha và cộng sự ghép đặc trưng của ResNet50 và MobileNetV2. Sau global average pooling, hai vector được chiếu về cùng 512 chiều, nối thành biểu diễn 1.024 chiều và đưa qua dense layers. Hybrid đạt accuracy 73,65% và macro F1 70,97% trên tập test độc lập của Soil_Data_V3.",
        "Điểm rút ra: feature fusion có thể khai thác đặc trưng bổ sung, nhưng kiến trúc hybrid phức tạp và tốn tài nguyên hơn một backbone nhẹ.",
        "Scientific Reports",
        "https://www.nature.com/articles/s41598-026-66201-0",
    ),
    (
        "5.5 Ensemble CNN và Vision Transformer",
        "Sheth và cộng sự kết hợp CNN, Vision Transformer, CycleGAN và fuzzy logic cho phân loại đất và gợi ý cây trồng. Nghiên cứu sử dụng dataset gốc 1.189 ảnh thuộc bảy loại đất và mở rộng dữ liệu bằng CycleGAN. Phương pháp hướng đến một hệ thống đầy đủ hơn nhưng có độ phức tạp cao.",
        "Điểm rút ra: ensemble và dữ liệu sinh là hướng mạnh khi có đủ tài nguyên; cần đánh giá cẩn thận để tránh hiệu năng cao do ảnh sinh quá giống dữ liệu huấn luyện.",
        "DOI",
        "https://doi.org/10.1016/j.engappai.2025.111425",
    ),
]

for heading, description, takeaway, label, url in papers:
    doc.add_paragraph(style="Heading 2").add_run(heading)
    body(doc, description)
    takeaway_paragraph = body(doc, takeaway, lead="Điểm rút ra:")
    takeaway_paragraph.paragraph_format.keep_with_next = True
    paper_link(doc, label, url)

doc.add_paragraph(style="Heading 1").add_run("6 So sánh các hướng giải pháp")

solutions = doc.add_table(rows=1, cols=5)
solutions.alignment = WD_TABLE_ALIGNMENT.CENTER
solutions.autofit = False
solution_widths = [Cm(3.0), Cm(3.2), Cm(3.2), Cm(3.0), Cm(3.5)]
for cell, text, width in zip(
    solutions.rows[0].cells,
    ["Hướng", "Dữ liệu nhỏ", "Chi phí", "Khả năng triển khai", "Nhận định từ literature"],
    solution_widths,
):
    cell.width = width
    cell.text = text
    shade(cell, "1F4E78")
    margins(cell)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)

solution_rows = [
    ("Đặc trưng thủ công + ML", "Khá phù hợp", "Thấp", "Dễ", "Dễ giải thích nhưng phụ thuộc thiết kế feature."),
    ("CNN từ đầu", "Rủi ro overfitting", "Trung bình", "Trung bình", "Cần đủ dữ liệu và ablation kiến trúc."),
    ("Transfer learning CNN nhẹ", "Phù hợp", "Thấp", "Rất tốt", "Cân bằng tốt giữa hiệu năng và tài nguyên."),
    ("Dense CNN", "Phù hợp", "Cao hơn", "Trung bình", "Có thể đạt accuracy cao nhưng mô hình nặng hơn."),
    ("Hybrid hoặc ensemble", "Cần kiểm soát", "Cao", "Khó hơn", "Đặc trưng phong phú nhưng phức tạp."),
    ("Tách nền + classifier", "Phù hợp", "Cao hơn", "Trung bình", "Giảm shortcut từ phông nền."),
]
for row_index, values in enumerate(solution_rows, start=1):
    cells = solutions.add_row().cells
    for col_index, (cell, value, width) in enumerate(zip(cells, values, solution_widths)):
        cell.width = width
        cell.text = value
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_index % 2 == 0:
            shade(cell, "EAF2F8")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index in (0, 4) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(8.3)
borders(solutions)

doc.add_paragraph(style="Heading 1").add_run("7 Phương pháp được lựa chọn từ literature")
body(
    doc,
    "Phương án được lựa chọn cho giai đoạn triển khai sau tuần đầu là transfer learning với MobileNetV3-Small, kết hợp chuẩn hóa ảnh, data augmentation và cơ chế xử lý class imbalance. Đây là quyết định được hình thành từ các kết quả và xu hướng trong paper, không phải từ kết quả thực nghiệm của dự án.",
)

doc.add_paragraph(style="Heading 2").add_run("7.1 Cơ sở chọn MobileNetV3 Small")
for text in (
    "Paper MobileNetV3 xây dựng phiên bản Small cho môi trường tài nguyên hạn chế và tối ưu độ trễ trên thiết bị di động.",
    "Light-SoilNet cho thấy lightweight CNN phù hợp với bài toán phân loại năm loại đất và dataset nhỏ.",
    "SoilNet-XAI cho thấy transfer learning giúp tận dụng backbone pretrained khi chỉ có vài trăm ảnh.",
    "Nghiên cứu so sánh transfer learning chỉ ra MobileNet có lợi thế rõ về tốc độ và bộ nhớ so với các mạng dày hơn.",
    "Các mô hình hybrid và ensemble có khả năng biểu diễn mạnh nhưng chưa phù hợp làm phương án đầu tiên khi tài nguyên và dữ liệu còn hạn chế.",
):
    bullet(doc, text)

doc.add_paragraph(style="Heading 2").add_run("7.2 Các thành phần phương pháp dự kiến")
for text in (
    "Đưa ảnh về kích thước chuẩn 224 x 224 và normalize theo backbone pretrained.",
    "Sử dụng data augmentation trên tập train để tăng biến thiên về góc nhìn, vị trí và ánh sáng.",
    "Dùng MobileNetV3-Small pretrained trên ImageNet làm bộ trích xuất đặc trưng.",
    "Thay lớp phân loại cuối theo số lớp của dataset mục tiêu.",
    "Dùng class-weighted loss hoặc sampling phù hợp nếu số ảnh giữa các lớp không cân bằng.",
    "Đánh giá bằng accuracy, precision, recall, macro F1 và confusion matrix.",
    "Ở bước phát triển sau, bổ sung Grad-CAM và kiểm thử ngoài miền để đánh giá độ tin cậy.",
):
    bullet(doc, text)

doc.add_paragraph(style="Heading 2").add_run("7.3 Vì sao chưa chọn các phương pháp khác")
body(
    doc,
    "DenseNet121 là đối chứng quan trọng vì đạt kết quả cao trong nhiều paper, nhưng có chi phí lớn hơn một mô hình mobile. Tách nền bằng YOLOv5 có giá trị khi background gây nhiễu, song làm pipeline phức tạp hơn. Hybrid CNN và Vision Transformer phù hợp khi có dataset lớn và tài nguyên huấn luyện tốt hơn. Vì vậy, MobileNetV3-Small được chọn làm phương án ban đầu để tạo một baseline nhẹ, có cơ sở khoa học và dễ mở rộng sang các thí nghiệm phức tạp hơn.",
)

doc.add_paragraph(style="Heading 1").add_run("8 Câu hỏi nghiên cứu cho giai đoạn tiếp theo")
for text in (
    "MobileNetV3-Small có đạt hiệu năng cạnh tranh trên dataset ảnh đất mục tiêu hay không?",
    "Data augmentation và class-weighted loss ảnh hưởng như thế nào đến các lớp ít ảnh?",
    "Mô hình nhẹ đánh đổi bao nhiêu accuracy so với DenseNet121 hoặc EfficientNet-B0?",
    "Confusion matrix cho thấy những lớp đất nào khó phân biệt nhất?",
    "Mô hình có duy trì confidence và tính nhất quán trên ảnh từ nguồn độc lập hay không?",
    "Grad-CAM có xác nhận rằng mô hình tập trung vào bề mặt đất thay vì phông nền không?",
):
    bullet(doc, text)

doc.add_paragraph(style="Heading 1").add_run("9 Kết luận tuần đầu")
body(
    doc,
    "Qua khảo sát literature, bài toán được xác định là supervised multiclass soil image classification với input là ảnh RGB và output là nhãn loại đất. Các metric cần theo dõi gồm accuracy, precision, recall, F1-score và confusion matrix. Transfer learning là hướng phù hợp khi dữ liệu hạn chế; CNN nhẹ hỗ trợ triển khai trên phần cứng phổ thông; augmentation và class balancing hỗ trợ giảm overfitting và mất cân bằng lớp.",
)
body(
    doc,
    "MobileNetV3-Small được chọn làm phương pháp dự kiến vì phù hợp với các kết luận từ survey, paper Light-SoilNet, SoilNet-XAI và nghiên cứu so sánh transfer learning. Đây là lựa chọn có chủ đích giữa hiệu năng, kích thước mô hình, tốc độ và khả năng triển khai. Kết quả thực nghiệm sẽ được trình bày ở các tuần sau, sau khi phương pháp đã được triển khai và đánh giá.",
)

doc.add_paragraph(style="Heading 1").add_run("10 Tài liệu tham khảo")
refs = [
    ("SoilClass: A Novel Image Dataset for Soil Classification. Mendeley Data, Version 1, 2026. DOI: 10.17632/6mx34nshm4.1.", "Mendeley Data", "https://data.mendeley.com/datasets/6mx34nshm4/1"),
    ("Howard, A. et al. Searching for MobileNetV3. Proceedings of ICCV, 2019, pp. 1314-1324.", "CVF Open Access", "https://openaccess.thecvf.com/content_ICCV_2019/html/Howard_Searching_for_MobileNetV3_ICCV_2019_paper.html"),
    ("Pandiri, D. N. K., Murugan, R., and Goel, T. Smart soil image classification system using lightweight convolutional neural network. Expert Systems with Applications, 238, 122185, 2024.", "DOI", "https://doi.org/10.1016/j.eswa.2023.122185"),
    ("SoilNet-XAI: A Deep Learning Approach with Explainability for Soil Image Classification. 28th ICCIT, IEEE, 2025. DOI: 10.1109/ICCIT68739.2025.11491058.", "IEEE Xplore", "https://ieeexplore.ieee.org/document/11491058/"),
    ("Cetin, N. et al. Performance comparison of deep and transfer learning models for smart soil texture classification. Computers and Electronics in Agriculture, 237, 110722, 2025.", "ScienceDirect", "https://www.sciencedirect.com/science/article/pii/S0168169925008282"),
    ("Sreelatha, A. et al. Hybrid deep learning model for automated soil classification and crop recommendation to enhance agricultural productivity. Scientific Reports, 2026.", "Nature", "https://www.nature.com/articles/s41598-026-66201-0"),
    ("Sheth, F. et al. An advanced artificial intelligence framework integrating ensembled CNNs and Vision Transformers for precise soil classification with adaptive fuzzy logic-based crop recommendations. Engineering Applications of Artificial Intelligence, 158, 111425, 2025.", "DOI", "https://doi.org/10.1016/j.engappai.2025.111425"),
    ("Soil image classification and segmentation: A survey from deep learning, multimodal data and hybrid models. GeoAI, 8, 100053, 2026.", "DOI", "https://doi.org/10.1016/j.geoai.2026.100053"),
]
for index, values in enumerate(refs, start=1):
    reference(doc, index, *values)

doc.core_properties.title = "Tổng quan bài toán phân loại ảnh đất"
doc.core_properties.subject = "Báo cáo đọc paper tuần đầu và định hướng lựa chọn phương pháp"
doc.core_properties.keywords = "soil classification, literature review, MobileNetV3, transfer learning"
doc.core_properties.author = "Nhóm dự án SoilClass"
doc.save(OUTPUT)
print(OUTPUT)
